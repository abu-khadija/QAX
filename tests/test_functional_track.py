"""
test_functional_track.py — End-to-end functional QA (the "third track")
=======================================================================
input → generate → plan → export → execute → report → results to Xray, and the
web faces of it. Regressions for the gaps found in the 2026-10-19 audit:

* a PR's fetched content was discarded (re-read from disk → 0 cases);
* `success` was True even when export or execution failed;
* execution results were never reported or sent to Xray — and the old upload
  posted Jira comments to LOCAL ids (TC-…) instead of the created Jira keys;
* `execute --suite` ran Jira issues that carried no steps;
* `run --mode smoke|regression` was a placeholder that executed nothing;
* `analyze` pushed live without asking when DRY_RUN=false;
* the web UI could not execute generated cases, skipped validation, and never
  sent the API token.

Offline: Xray/Jira/HTTP/Claude are all stubbed; nothing leaves the machine.
"""

import json

import pytest

# ── Xray client: results import + steps fetch ────────────────────────────────


class _Resp:
    """Minimal requests.Response stand-in."""

    def __init__(self, status=200, payload=None, text=""):
        self.status_code = status
        self._payload = payload
        self.text = text or json.dumps(payload or {})
        self.content = self.text.encode()

    def json(self):
        return self._payload

    def raise_for_status(self):
        if self.status_code >= 400:
            import requests
            err = requests.HTTPError(f"{self.status_code}")
            err.response = self
            raise err


@pytest.fixture
def xray_env(monkeypatch):
    """Fake Xray/Jira credentials (never contacted: requests is patched in each test)."""
    for k, v in {"XRAY_CLIENT_ID": "id", "XRAY_CLIENT_SECRET": "secret", "JIRA_CLOUD_ID": "cloud",
                 "JIRA_BASE_URL": "https://jira.invalid", "JIRA_EMAIL": "qa@invalid",
                 "JIRA_API_TOKEN": "tok", "JIRA_PROJECT_KEY": "QA"}.items():
        monkeypatch.setenv(k, v)


def test_import_results_dry_run_sends_nothing_and_shows_the_payload(xray_env, monkeypatch):
    import core.xray_client as xc
    monkeypatch.setattr(xc.requests, "post", lambda *a, **k: pytest.fail("network in dry run"))
    out = xc.XrayClient(dry_run=True).import_execution_results(
        [{"testKey": "QA-1", "status": "PASSED"}], summary="run")
    assert out["success"] and out["dry_run"] and out["key"] is None
    assert out["payload"]["tests"][0]["testKey"] == "QA-1"
    assert out["payload"]["info"]["project"] == "QA"


def test_import_results_live_posts_to_the_xray_execution_import(xray_env, monkeypatch):
    import core.xray_client as xc
    calls = []

    def post(url, **kw):
        calls.append((url, kw.get("json")))
        if url.endswith("/authenticate"):
            return _Resp(200, None, text='"TOKEN"')
        return _Resp(200, {"id": "1", "key": "QA-99", "self": "x"})
    monkeypatch.setattr(xc.requests, "post", post)
    out = xc.XrayClient(dry_run=False).import_execution_results(
        [{"testKey": "QA-1", "status": "FAILED"}])
    assert out["success"] and out["key"] == "QA-99"
    assert calls[-1][0].endswith("/api/v2/import/execution")
    assert calls[-1][1]["tests"] == [{"testKey": "QA-1", "status": "FAILED"}]


def test_import_results_reports_http_failure_and_empty_input(xray_env, monkeypatch):
    import core.xray_client as xc

    def post(url, **kw):
        if url.endswith("/authenticate"):
            return _Resp(200, None, text='"TOKEN"')
        return _Resp(400, {"error": "bad"}, text="no results")
    monkeypatch.setattr(xc.requests, "post", post)
    c = xc.XrayClient(dry_run=False)
    out = c.import_execution_results([{"testKey": "QA-1", "status": "PASSED"}])
    assert not out["success"] and "HTTP 400" in out["error"]
    assert c.import_execution_results([])["error"]           # nothing to send ≠ success


def test_fetch_test_steps_parses_graphql_and_is_silent_in_dry_run(xray_env, monkeypatch):
    import core.xray_client as xc
    assert xc.XrayClient(dry_run=True).fetch_test_steps(["QA-1"]) == {}

    def post(url, **kw):
        if url.endswith("/authenticate"):
            return _Resp(200, None, text='"TOKEN"')
        assert url.endswith("/graphql") and 'key in ("QA-1","QA-2")' in kw["json"]["variables"]["jql"]
        return _Resp(200, {"data": {"getTests": {"results": [
            {"jira": {"key": "QA-1"}, "steps": [{"action": "Open /login", "data": "", "result": "Form shown"}]},
            {"jira": {"key": "QA-2"}, "steps": []}]}}})
    monkeypatch.setattr(xc.requests, "post", post)
    steps = xc.XrayClient(dry_run=False).fetch_test_steps(["QA-1", "QA-2", "DRY-RUN-X"])
    assert steps == {"QA-1": [{"action": "Open /login", "data": "", "expected_result": "Form shown"}],
                     "QA-2": []}


def test_fetch_tests_by_tag_falls_back_when_the_new_search_endpoint_is_missing(xray_env, monkeypatch):
    import core.xray_client as xc
    seen = []

    def get(url, **kw):
        seen.append(url)
        if url.endswith("/search/jql"):
            return _Resp(410, {}, text="gone")
        return _Resp(200, {"issues": [{"key": "QA-7", "fields": {"summary": "s", "labels": ["smoke"]}}]})
    monkeypatch.setattr(xc.requests, "get", get)
    tests = xc.XrayClient(dry_run=False).fetch_tests_by_tag("smoke")
    assert [t["key"] for t in tests] == ["QA-7"] and seen[-1].endswith("/rest/api/3/search")


def test_issues_to_cases_keeps_the_jira_key_and_steps():
    from core.xray_client import issues_to_cases
    cases = issues_to_cases([{"key": "QA-1", "summary": "Login"}, {"key": "QA-2", "summary": "x"}],
                            {"QA-1": [{"action": "a", "data": "", "expected_result": "r"}]})
    assert cases[0]["test_case_id"] == cases[0]["jira_key"] == "QA-1"
    assert cases[0]["manual_steps"][0] == {"step_number": 1, "action": "a", "data": "",
                                           "expected_result": "r"}
    assert cases[1]["manual_steps"] == []                    # → reported as not run, never PASS


# ── Results → Xray attribution ───────────────────────────────────────────────


class _FakeXray:
    """Records what report_to_xray asks Xray to import."""
    XRAY_STATUS = {"PASS": "PASSED", "HEALED": "PASSED", "FAIL": "FAILED", "SKIP": "TODO", "ERROR": "TODO"}

    def __init__(self, dry_run=True):
        self.dry_run = dry_run
        self.sent = None

    def import_execution_results(self, tests, summary=""):
        self.sent = tests
        return {"success": True, "dry_run": self.dry_run, "key": None if self.dry_run else "QA-50"}


def _run_report(batch, client, **kw):
    import asyncio

    from execution.results_reporter import ResultsReporter
    return asyncio.run(ResultsReporter(dry_run=client.dry_run).report_to_xray(batch, client, **kw))


def test_results_never_use_local_ids_as_jira_keys(monkeypatch):
    monkeypatch.setenv("DRY_RUN", "false")
    client = _FakeXray(dry_run=False)
    batch = {"results": [{"test_case_id": "TC-1", "status": "PASS"},          # looks like a key, is not
                         {"test_case_id": "TC-AUTH-0002", "status": "SKIP"},
                         {"test_case_id": "QA-3", "status": "FAIL", "jira_key": "QA-3"}]}
    out = _run_report(batch, client, key_map={"TC-AUTH-0002": "QA-2"})
    assert [t["testKey"] for t in client.sent] == ["QA-2", "QA-3"]   # TC-1 is never guessed
    assert out["unmapped"] == 1 and out["uploaded"] == 2 and out["test_execution_key"] == "QA-50"
    by_key = {t["testKey"]: t for t in client.sent}
    assert by_key["QA-2"]["status"] == "TODO" and "NOT a pass" in by_key["QA-2"]["comment"]
    assert by_key["QA-3"]["status"] == "FAILED"


def test_results_with_no_jira_key_send_nothing_and_say_why():
    client = _FakeXray(dry_run=True)
    out = _run_report({"results": [{"test_case_id": "TC-1", "status": "PASS"}]}, client)
    assert client.sent is None and out["uploaded"] == 0 and "no result has a Jira test key" in out["error"]


def test_xray_adapter_key_map_excludes_dry_run_placeholders(monkeypatch):
    import core.xray_client as xc
    from core.adapters.xray_adapter import XrayAdapter

    def push_batch(self, cases):
        return {"pushed": [{"test_case_id": "TC-1", "jira_key": "QA-1"},
                           {"test_case_id": "TC-2", "jira_key": "DRY-RUN-TC-2"}], "failed": []}
    monkeypatch.setattr(xc.XrayClient, "push_batch", push_batch)
    out = XrayAdapter().export([{"test_case_id": "TC-1"}, {"test_case_id": "TC-2"}], dry_run=True)
    assert out["key_map"] == {"TC-1": "QA-1"}


# ── Pipeline ─────────────────────────────────────────────────────────────────


class _Adapter:
    """Export adapter stub."""

    def __init__(self, result):
        self.result = result

    def export(self, cases, project=None, dry_run=True):
        return dict(self.result)


def _pipeline(tmp_path, *, export="none", exec_mode="none", adapter=None, execution=None,
              dry_run=True, base_url="http://app.local"):
    from core.pipeline import QAXPipeline
    p = QAXPipeline(export=export, exec_mode=exec_mode, dry_run=dry_run,
                    base_url=base_url, output_dir=str(tmp_path / "run"))
    if adapter is not None:
        p.adapter = adapter
    if execution is not None:
        p._run_execution = lambda cases: execution
    return p


_CASES = [{"test_case_id": "TC-1", "title": "a", "api_steps": [{"method": "GET", "path": "/"}]},
          {"test_case_id": "TC-2", "title": "b", "manual_steps": [{"action": "open"}]}]


def test_plan_describes_runnable_cases_writes_and_missing_prerequisites(tmp_path):
    p = _pipeline(tmp_path, export="xray", exec_mode="api", base_url="")
    plan = p.plan(_CASES)
    assert plan["cases"] == 2 and plan["api_executable"] == 1 and plan["ui_executable"] == 1
    assert plan["live_writes"] is False and any("dry run" in w for w in plan["writes"])
    assert any("no target URL" in m for m in plan["missing"])


def test_plan_only_stops_before_export_and_execution(tmp_path):
    p = _pipeline(tmp_path, exec_mode="api",
                  adapter=_Adapter({"success": True}),
                  execution={"error": "must not run"})
    p.adapter.export = lambda *a, **k: pytest.fail("exported in plan-only")
    p._run_execution = lambda *a: pytest.fail("executed in plan-only")
    r = p.run(test_cases=[dict(c) for c in _CASES], plan_only=True)
    assert r["success"] and r["plan"]["cases"] == 2 and r["export"] is None


def test_success_is_false_when_export_or_execution_fails(tmp_path):
    p = _pipeline(tmp_path, export="xray", adapter=_Adapter({"success": False, "error": "401"}))
    r = p.run(test_cases=[dict(c) for c in _CASES])
    assert r["success"] is False and any("export" in e for e in r["errors"])

    p = _pipeline(tmp_path, exec_mode="api", adapter=_Adapter({"success": True}),
                  execution={"error": "connection refused", "total": 2, "results": []})
    r = p.run(test_cases=[dict(c) for c in _CASES])
    assert r["success"] is False and any("connection refused" in e for e in r["errors"])


def test_export_keys_are_stamped_report_written_and_results_attributed(tmp_path, monkeypatch, xray_env):
    import core.xray_client as xc
    sent = {}

    def fake_import(self, tests, summary="", test_execution_key=None):
        sent["tests"] = tests
        return {"success": True, "dry_run": self.dry_run, "key": "QA-50", "payload": {}}
    monkeypatch.setattr(xc.XrayClient, "import_execution_results", fake_import)
    monkeypatch.setenv("DRY_RUN", "false")
    batch = {"total": 2, "passed": 1, "failed": 0, "skipped": 1, "mode": "api",
             "results": [{"test_case_id": "TC-1", "status": "PASS", "steps": []},
                         {"test_case_id": "TC-2", "status": "SKIP", "steps": []}]}
    p = _pipeline(tmp_path, export="xray", exec_mode="api", dry_run=False,
                  adapter=_Adapter({"success": True, "exported_count": 2, "references": ["QA-1", "QA-2"],
                                    "key_map": {"TC-1": "QA-1", "TC-2": "QA-2"}}),
                  execution=batch)
    cases = [dict(c) for c in _CASES]
    r = p.run(test_cases=cases)
    assert [c["jira_key"] for c in cases] == ["QA-1", "QA-2"]
    assert {t["testKey"]: t["status"] for t in sent["tests"]} == {"QA-1": "PASSED", "QA-2": "TODO"}
    assert r["success"] and r["results_upload"][0]["test_execution_key"] == "QA-50"
    saved = json.loads((tmp_path / "run" / "test_cases.json").read_text())
    assert saved["test_cases"][0]["jira_key"] == "QA-1"
    assert (tmp_path / "run" / "results.json").exists() and r["report"]["html"]


def test_dry_run_results_upload_explains_missing_keys_without_failing(tmp_path, monkeypatch):
    monkeypatch.setenv("DRY_RUN", "true")
    batch = {"total": 1, "passed": 1, "mode": "api",
             "results": [{"test_case_id": "TC-1", "status": "PASS", "steps": []}]}
    p = _pipeline(tmp_path, export="xray", exec_mode="api",
                  adapter=_Adapter({"success": True, "exported_count": 2, "key_map": {}}),
                  execution=batch)
    r = p.run(test_cases=[dict(c) for c in _CASES])
    assert r["success"], r["errors"]
    assert "dry run" in r["results_upload"][0]["error"]


def test_pr_input_uses_the_fetched_content_not_the_local_disk(tmp_path):
    p = _pipeline(tmp_path)
    seen = {}

    def gen(file_path, **kw):
        seen.update(kw, file_path=file_path)
        return []
    p._run_generation = gen
    p._run_generation_from_item({"source_type": "pr", "file_path": "src/not_checked_out.py",
                                 "content": "def f(): pass", "metadata": {}})
    assert seen["content"] == "def f(): pass" and seen["file_path"] == "src/not_checked_out.py"


def test_file_input_still_reads_from_disk(tmp_path):
    src = tmp_path / "a.py"
    src.write_text("x = 1")
    p = _pipeline(tmp_path)
    seen = {}
    p._run_generation = lambda file_path, **kw: seen.update(kw) or []
    p._run_generation_from_item({"source_type": "file", "file_path": str(src),
                                 "content": "x = 1", "metadata": {}})
    assert "content" not in seen


# ── CLI: analyze live push needs a yes ──────────────────────────────────────


def test_live_push_needs_yes_or_an_interactive_y(monkeypatch):
    import main
    assert main._confirm_live_push(3, True) is True
    monkeypatch.setattr(main.sys.stdin, "isatty", lambda: False, raising=False)
    assert main._confirm_live_push(3, False) is False          # CI without --yes never writes


def test_pipeline_cli_is_dry_unless_live_and_env_allow(monkeypatch):
    import main
    captured = {}

    class _P:
        def __init__(self, **kw):
            captured.update(kw)
            self.project = ""

        def run(self, **kw):
            return {"generation": {"count": 0, "success": True}, "success": True, "errors": []}
    monkeypatch.setattr("core.pipeline.QAXPipeline", _P)
    for env, flags, expect_dry in (("false", [], True), ("false", ["--live"], False),
                                   ("true", ["--live"], True)):
        monkeypatch.setenv("DRY_RUN", env)
        monkeypatch.setattr(main.sys, "argv", ["main.py", "pipeline", "--input", "story",
                                               "--source", "s", "--base-url", "http://a"] + flags)
        main.run_pipeline_command()
        assert captured["dry_run"] is expect_dry and captured["base_url"] == "http://a"


# ── run --mode smoke|regression executes ─────────────────────────────────────


def test_smoke_run_executes_fetched_tests(monkeypatch):
    import core.xray_client as xc
    import services.execution_service as es
    from core.session_manager import SessionManager
    from services.results import ExecutionResult

    monkeypatch.setattr(xc.XrayClient, "fetch_test_steps",
                        lambda self, keys: {"QA-1": [{"action": "Open /", "data": "", "expected_result": "ok"}]})
    ran = {}

    def execute(self, cases, **kw):
        ran["cases"] = cases
        return ExecutionResult(mode="ui", raw={"total": 1, "passed": 1, "failed": 0, "results": []})
    monkeypatch.setattr(es.ExecutionService, "execute", execute)
    sm = SessionManager.__new__(SessionManager)
    out = sm._execute_fetched_tests([{"key": "QA-1", "summary": "Home"}], {"headless": "true"}, 0.0)
    assert ran["cases"][0]["jira_key"] == "QA-1" and out["executed"] == 1 and out["passed"] == 1


def test_smoke_run_without_xray_credentials_reports_not_run_instead_of_crashing(monkeypatch):
    import services.execution_service as es
    from core.session_manager import SessionManager
    from services.results import ExecutionResult
    for k in ("XRAY_CLIENT_ID", "XRAY_CLIENT_SECRET", "JIRA_CLOUD_ID"):
        monkeypatch.setenv(k, "")
    ran = {}

    def execute(self, cases, **kw):
        ran["cases"] = cases
        return ExecutionResult(mode="ui", raw={"total": 1, "errors": 1, "results": []})
    monkeypatch.setattr(es.ExecutionService, "execute", execute)
    sm = SessionManager.__new__(SessionManager)
    out = sm._execute_fetched_tests([{"key": "QA-1", "summary": "Home"}], {}, 0.0)
    assert ran["cases"][0]["manual_steps"] == [] and out["passed"] == 0 and out["errors"] == 1


def test_pipeline_upload_without_xray_credentials_is_an_error_not_a_crash(tmp_path, monkeypatch):
    for k in ("XRAY_CLIENT_ID", "XRAY_CLIENT_SECRET", "JIRA_CLOUD_ID"):
        monkeypatch.setenv(k, "")
    monkeypatch.setenv("DRY_RUN", "false")
    batch = {"total": 1, "mode": "api", "results": [{"test_case_id": "TC-1", "status": "PASS"}]}
    p = _pipeline(tmp_path, export="xray", exec_mode="api", dry_run=False,
                  adapter=_Adapter({"success": True, "key_map": {"TC-1": "QA-1"}}), execution=batch)
    r = p.run(test_cases=[dict(c) for c in _CASES])
    assert r["success"] is False and any("Missing required" in e for e in r["errors"])


# ── Web faces ────────────────────────────────────────────────────────────────


@pytest.fixture
def client(monkeypatch, tmp_path):
    monkeypatch.chdir(tmp_path)                      # run history / outputs stay out of the repo
    monkeypatch.delenv("QAX_API_TOKEN", raising=False)
    monkeypatch.delenv("QAX_ALLOW_LIVE_EXPORT", raising=False)
    from web.server import app
    return app.test_client()


_WEB_CASES = [
    {"title": "list items", "api_steps": [{"method": "GET", "path": "/items", "expected_status": 200}]},
    {"title": "create item", "api_steps": [{"method": "POST", "path": "/items", "expected_status": 201}]},
    {"title": "ui only", "manual_steps": [{"action": "click save"}]},
]


def test_execute_preview_lists_what_runs_and_what_does_not(client):
    r = client.post("/api/execute", json={"tests": _WEB_CASES, "url": "http://app.local",
                                          "preview": True, "read_only": True})
    plan = r.get_json()["plan"]
    assert [x["title"] for x in plan["runnable"]] == ["list items"]
    reasons = {n["title"]: n["reason"] for n in plan["not_run"]}
    assert "read-only" in reasons["create item"] and "Browser engine" in reasons["ui only"]


def test_execute_requires_authorization_and_a_url(client):
    assert client.post("/api/execute", json={"tests": _WEB_CASES, "url": "app"}).status_code == 400
    r = client.post("/api/execute", json={"tests": _WEB_CASES, "url": "http://app.local"})
    assert r.status_code == 400 and "authorized" in r.get_json()["error"]


def test_execute_runs_runnable_cases_over_http(client, monkeypatch):
    import services.execution_service as es
    from services.results import ExecutionResult
    seen = {}

    def execute(self, cases, **kw):
        seen.update(kw, cases=cases)
        return ExecutionResult(mode="api", raw={"total": 2, "passed": 1, "failed": 1, "results": [
            {"test_case_id": c["test_case_id"], "status": "PASS"} for c in cases]})
    monkeypatch.setattr(es.ExecutionService, "execute", execute)
    r = client.post("/api/execute", json={"tests": _WEB_CASES, "url": "http://app.local",
                                          "authorized": True, "read_only": False, "token": "T"})
    d = r.get_json()
    assert r.status_code == 200 and d["passed"] == 1 and d["not_run"] == 1
    assert seen["mode"] == "api" and seen["base_url"] == "http://app.local"
    assert seen["auth"] == {"token": "T"} and len(seen["cases"]) == 2


def test_generate_goes_through_the_validating_generation_service(client, monkeypatch):
    import services.generation_service as gs
    from services.results import GenerationResult
    seen = {}

    def generate(self, source, options=None, progress=None):
        seen.update(source=source, content=options.content, story=options.user_story)
        return GenerationResult(source=source, test_cases=[{"title": "ok"}],
                                invalid=[{"test_case": {}, "error": "bad"}])
    monkeypatch.setattr(gs.GenerationService, "generate", generate)
    monkeypatch.setattr("services.context.QAXContext.require_anthropic_key", lambda self: None,
                        raising=False)
    d = client.post("/api/generate", json={"story": "As a user I can log in"}).get_json()
    assert d["tests"] == [{"title": "ok"}] and d["invalid"] == 1
    assert seen["source"].startswith("story://") and seen["content"] == "As a user I can log in"


def test_ui_config_and_page_wire_execution_token_and_honest_export(client, monkeypatch):
    cfg = client.get("/api/config").get_json()
    assert cfg["live_export"] is False and cfg["token_required"] is False and "browser" in cfg
    html = "".join(client.get(u).get_data(as_text=True) for u in ("/app", "/app.js", "/app.css"))
    for needle in ("/api/execute", "runTests(", "X-QAX-Token", "AI proposal · not run",
                   "QAX_ALLOW_LIVE_EXPORT", "dry_run:!live", "tc.manual_steps",
                   'id="exengine"', "loadTests(", "engine:$('exengine').value"):
        assert needle in html, needle
    assert "Add your Jira/Xray credentials to <code>.env</code> to file for real" not in html


def test_api_token_gate_covers_the_new_routes(monkeypatch):
    monkeypatch.setenv("QAX_API_TOKEN", "sekret")
    from web.server import app
    c = app.test_client()
    assert c.get("/api/config").status_code == 401
    assert c.post("/api/execute", json={}).status_code == 401
    assert c.get("/api/config", headers={"X-QAX-Token": "sekret"}).status_code == 200


def test_browser_engine_never_runs_in_read_only_and_needs_steps(client):
    cases = [{"title": "ui", "automation_steps": [{"playwright_action": "navigate", "value": "/"}]},
             {"title": "empty"}]
    plan = client.post("/api/execute", json={"tests": cases, "url": "http://app.local", "engine": "ui",
                                             "preview": True, "read_only": True}).get_json()["plan"]
    assert plan["engine"] == "ui" and plan["runnable"] == []
    assert {n["title"]: n["reason"] for n in plan["not_run"]} == {
        "ui": "read-only is on: the browser engine clicks and submits forms — untick read-only to run it (staging only)", "empty": "no steps"}
    plan = client.post("/api/execute", json={"tests": cases, "url": "http://app.local", "engine": "ui",
                                             "preview": True, "read_only": False}).get_json()["plan"]
    assert [x["title"] for x in plan["runnable"]] == ["ui"]


def test_web_browser_run_restores_env_and_forces_ai_healing_off(client, monkeypatch):
    import os

    import services.execution_service as es
    from services.results import ExecutionResult
    seen = {}

    def execute(self, cases, **kw):
        seen.update(kw, heal=os.environ.get("QAX_AI_HEAL"), base=os.environ.get("APP_BASE_URL"),
                    proj=os.environ.get("DEFAULT_PROJECT"))
        return ExecutionResult(mode="ui", raw={"total": 1, "passed": 1, "results": []})
    monkeypatch.setattr(es.ExecutionService, "execute", execute)
    monkeypatch.setenv("QAX_AI_HEAL", "true")
    monkeypatch.setenv("APP_BASE_URL", "http://before")
    r = client.post("/api/execute", json={
        "tests": [{"title": "t", "manual_steps": [{"action": "open /"}]}], "url": "http://app.local",
        "engine": "ui", "project": "uidemo/../x", "read_only": False, "authorized": True})
    assert r.status_code == 200
    assert seen["mode"] == "ui" and seen["heal"] == "false" and seen["base"] == "http://app.local"
    assert seen["proj"] == "uidemox"                                  # path characters stripped
    assert os.environ["QAX_AI_HEAL"] == "true" and os.environ["APP_BASE_URL"] == "http://before"


def test_ai_healing_is_off_unless_opted_in(monkeypatch):
    from execution.test_executor import ai_assist_enabled
    monkeypatch.delenv("QAX_AI_HEAL", raising=False)
    monkeypatch.setenv("ANTHROPIC_API_KEY", "sk-test")
    assert ai_assist_enabled() is False
    monkeypatch.setenv("QAX_AI_HEAL", "true")
    assert ai_assist_enabled() is True


def test_deterministic_labels_name_fields_buttons_and_links_without_ai():
    from core.dom_analyzer import label_elements
    sel = label_elements([
        {"tag": "input", "id": "id_username", "name": "username", "type": "text"},          # Django ids
        {"tag": "input", "id": "", "name": "", "type": "submit", "value": "Login"},
        {"tag": "a", "href": "/bookmarks/new", "text": "Add bookmark"},
        {"tag": "a", "href": "", "aria": "Add bookmark", "text": ""},                     # same name
        {"tag": "button", "text": "Logout"},
        {"tag": "input", "id": "q", "name": "q", "type": "search", "label": "Search"},
        {"tag": "select", "name": "sort"},
        {"tag": "div"},                                                                      # nameless
    ])
    assert sel["username"] == "#id_username"
    assert sel["login_button"] == "input[type='submit'][value='Login']"
    assert sel["add_bookmark_link"] == "a[href='/bookmarks/new']"
    assert sel["logout_button"] == "button:has-text('Logout')"
    assert sel["search"] == "#q" and sel["sort_select"] == "select[name='sort']"
    assert sel["add_bookmark_link_2"] == "a[aria-label='Add bookmark']"   # 2nd element, own key
    assert len(sel) == 7                                     # the nameless div is skipped


def test_env_file_fills_gaps_but_never_overrides_what_the_operator_set(tmp_path, monkeypatch):
    """A .env with DRY_RUN=false used to override DRY_RUN=true set on the command line / in CI."""
    from core.env_loader import load_env
    env = tmp_path / ".env"
    env.write_text("DRY_RUN=false\nQAX_T_KEY=from-file\nQAX_T_EMPTY=filled\n")
    monkeypatch.setenv("DRY_RUN", "true")
    monkeypatch.delenv("QAX_T_KEY", raising=False)
    monkeypatch.setenv("QAX_T_EMPTY", "")
    applied = load_env(str(env))
    import os
    assert os.environ["DRY_RUN"] == "true"                      # the operator's choice wins
    assert os.environ["QAX_T_KEY"] == "from-file" and os.environ["QAX_T_EMPTY"] == "filled"
    assert "DRY_RUN" not in applied


def test_the_app_has_two_tracks_and_generation_lives_inside_functional_qa(client):
    html = "".join(client.get(u).get_data(as_text=True) for u in ("/app", "/app.js", "/app.css"))
    assert 'id="mode-scan"' in html and 'id="mode-func"' in html and 'id="mode-gen"' not in html
    assert 'id="panel-gen"' not in html and 'id="src-gen"' not in html
    assert "if(m==='gen'){ m='func'; }" in html                                         # old links
    # one E2E flow: inputs → understand → cases (export/share) → run → report → bugs to Jira
    for needle in ('id="mapurl"', 'id="fqrepo"', 'id="story"', 'analyzeAll()', "/api/functional/analyze",
                   "downloadCases('csv')", "Share to Jira / Xray", "renderReport(d)", "/api/functional/bugs",
                   "Functional QA (E2E)"):
        assert needle in html, needle
    assert "Generate test cases from a story" not in html
    landing = client.get("/").get_data(as_text=True)
    assert "/app#functional" in landing and "/app#generate" not in landing


def test_generation_errors_are_explained_in_plain_words(client, monkeypatch):
    import services.generation_service as gs
    from services.results import GenerationResult

    def generate(self, source, options=None, progress=None):
        return GenerationResult(source=source, error="Error code: 400 - {'message': 'Your credit balance "
                                                     "is too low to access the Anthropic API.'}")
    monkeypatch.setattr(gs.GenerationService, "generate", generate)
    d = client.post("/api/generate", json={"story": "s"}).get_json()
    assert "no credit left" in d["error"] and "Load a test-cases file" in d["error"]
    assert "{'message'" not in d["error"]


def test_page_templates_and_risky_links():
    from core.app_mapper import is_risky_link, page_template
    assert page_template("http://a/bookmarks/12/edit?return_url=/x") == "/bookmarks/{id}/edit"
    assert page_template("http://a/bookmarks?page=2") == "/bookmarks"
    assert page_template("http://a/u/3f2a9c1e-77aa-4b2c/x") == "/u/{id}/x"
    for h in ("/logout/", "/settings/integrations/create-api-token", "/items/3/delete", "/export.csv"):
        assert is_risky_link(h), h
    for h in ("/bookmarks/archived", "/bookmarks/new", "/tags/merge", "/exports", "/runs"):
        assert not is_risky_link(h), h


def _a_map():
    from core.app_mapper import AppMap
    return AppMap(base_url="http://a", pages=[
        {"path": "/login", "elements": {"username": "#u", "login_button": "#go"}},
        {"path": "/notes", "elements": {"note_input": "#n"}}])


def test_check_cases_flags_steps_on_elements_that_are_not_on_the_map():
    from core.app_mapper import check_cases
    cases = [
        {"automation_steps": [{"playwright_action": "fill", "selector_key": "username"},
                              {"playwright_action": "assert", "selector_key": "page"},
                              {"playwright_action": "click", "selector_key": "#go"}]},
        {"automation_steps": [{"playwright_action": "navigate", "selector_key": "TODO_url"},
                              {"playwright_action": "click", "selector_key": "save_btn"}]}]
    assert check_cases(cases, _a_map()) == 1
    assert cases[0]["unmapped_elements"] == [] and "needs_review" not in cases[0]
    assert cases[1]["unmapped_elements"] == ["save_btn"] and cases[1]["needs_review"] is True
    assert check_cases(cases, None) == 0


def test_generation_is_grounded_in_the_map_and_flags_unknown_elements(monkeypatch):
    import core.claude_client as cc
    from services.context import QAXContext
    from services.generation_service import GenerationOptions, GenerationService
    seen = {}

    def analyze_code(self, *a, **kw):
        seen["ctx"] = kw.get("repo_context") or ""
        seen["cache"] = kw.get("use_cache")
        from tests.test_validator import make_valid_test_case
        tc = make_valid_test_case()
        steps = tc["automation_steps"]
        steps[0]["selector_key"] = "username"                 # on the map
        steps[-1]["selector_key"] = "ghost_button"            # not on the map
        steps[-1]["playwright_action"] = "click"
        return [tc]
    monkeypatch.setattr(cc.ClaudeClient, "__init__", lambda self, *a, **k: None)
    monkeypatch.setattr(cc.ClaudeClient, "analyze_code", analyze_code)
    monkeypatch.setattr(QAXContext, "require_anthropic_key", lambda self: None, raising=False)
    res = GenerationService(QAXContext.from_env()).generate(
        "story://t", options=GenerationOptions(content="As a user I log in", user_story="log in",
                                               app_map=_a_map()))
    assert "OBSERVED UI SURFACE" in seen["ctx"] and "username: #u" in seen["ctx"]
    assert seen["cache"] is False
    assert len(res.test_cases) == 1, res.invalid            # the case passed validation …
    assert res.test_cases[0]["unmapped_elements"] == ["ghost_button"]    # … and was flagged
    assert res.summary["unmapped_cases"] == 1


def test_web_map_needs_authorization_and_saves_the_map(client, monkeypatch, tmp_path):
    import core.app_mapper as am
    monkeypatch.setattr(am, "CONFIG_DIR", tmp_path)

    async def fake_map(url, **kw):
        assert kw["username"] == "u" and kw["password"] == "p"
        return _a_map()
    monkeypatch.setattr(am, "map_app", fake_map)
    body = {"url": "http://app.local", "project": "my/../app", "username": "u", "password": "p"}
    assert client.post("/api/map", json=body).status_code == 400            # not authorized
    d = client.post("/api/map", json=dict(body, authorized=True)).get_json()
    assert d["project"] == "myapp" and [p["path"] for p in d["pages"]] == ["/login", "/notes"]
    saved = (tmp_path / "myapp" / "app_map.json").read_text()
    assert '"p"' not in saved and "password" not in saved.lower()          # login never stored
    assert (tmp_path / "myapp" / "selectors.json").exists()


def test_pipeline_plan_says_when_there_is_no_app_map(tmp_path):
    p = _pipeline(tmp_path, exec_mode="ui")
    p.app_map = None
    assert any("no app map" in m for m in p.plan([])["missing"])
    p.app_map = _a_map()
    plan = p.plan([])
    assert plan["app_map_pages"] == 2 and not any("no app map" in m for m in plan["missing"])


def test_artifact_route_serves_only_run_screenshots_and_traces(client, tmp_path):
    from pathlib import Path
    root = Path(__file__).resolve().parent.parent
    shots = root / "output" / "screenshots" / "steps"
    shots.mkdir(parents=True, exist_ok=True)
    ok = shots / "qax_test_artifact.png"
    ok.write_bytes(b"\x89PNG\r\n\x1a\n")
    try:
        assert client.get("/api/artifact?path=output/screenshots/steps/qax_test_artifact.png").status_code == 200
        for bad in ("output/screenshots/../../.env", ".env", "/etc/passwd", "README.md",
                    "output/screenshots/steps/../../../main.py"):
            assert client.get("/api/artifact", query_string={"path": bad}).status_code == 404, bad
    finally:
        ok.unlink()


# ── E2E flow: understand → trace → bugs ──────────────────────────────────────

def test_acceptance_criteria_are_numbered_and_coverage_is_traced():
    from core.traceability import coverage, criteria_block, extract_criteria
    story = "As a user I can reset my password.\n- The link expires in 15 minutes\n- A used link shows an error"
    ac = extract_criteria(story)
    assert [c["id"] for c in ac] == ["AC-1", "AC-2"] and ac[0]["text"] == "The link expires in 15 minutes"
    assert [c["id"] for c in extract_criteria("Users can log in. Wrong passwords are rejected.")] == ["AC-1", "AC-2"]
    assert extract_criteria("") == [] and criteria_block([]) == ""
    assert "AC-2: A used link shows an error" in criteria_block(ac)
    cases = [{"test_case_id": "T1", "description": "Covers: AC-1 | Technique: Boundary Value Analysis. x"},
             {"test_case_id": "T2", "description": "Covers: AC-1, AC-9 | Technique: Error Guessing."},
             {"test_case_id": "T3", "description": "no trace line"}]
    cov = coverage(cases, ac)
    assert cov["by_criterion"] == {"AC-1": ["T1", "T2"], "AC-2": []}
    assert cov["uncovered"] == ["AC-2"] and cov["untraced_cases"] == ["T3"] and cov["percent"] == 50
    assert cov["techniques"] == {"Boundary Value Analysis": 1, "Error Guessing": 1}


def test_a_failed_e2e_case_becomes_an_istqb_shaped_bug():
    from core.jira_bugs import functional_result_to_bug
    case = {"test_case_id": "T1", "title": "Reset link expires", "priority": "High",
            "description": "Covers: AC-1 | Technique: BVA.",
            "manual_steps": [{"step_number": 1, "action": "Request a reset link"},
                             {"step_number": 2, "action": "Open it after 16 minutes", "expected_result": "Link expired"}]}
    result = {"test_case_id": "T1", "status": "FAIL",
              "steps": [{"step_number": 1, "status": "PASS"},
                        {"step_number": 2, "status": "FAIL", "actual_result": "Password form shown",
                         "screenshot": "output/screenshots/t1.png"}],
              "page_errors": [{"kind": "HTTP 500", "text": "POST /reset"}]}
    bug = functional_result_to_bug(result, case, base_url="http://app.local")
    assert bug["title"] == "[P2] Reset link expires" and bug["severity"] == "P2"
    d = bug["description"]
    assert "Requirement: AC-1" in d and "1. Open http://app.local — Request a reset link" in d
    assert "Expected: Link expired" in d and "Actual: Password form shown" in d
    assert "HTTP 500: POST /reset" in d and "output/screenshots/t1.png" in d
    assert "qax" not in functional_result_to_bug(result, case, client_safe=True)["labels"]


def test_web_functional_analyze_reads_understands_and_traces(client, monkeypatch, tmp_path):
    import core.app_mapper as am
    import services.generation_service as gs
    from services.results import GenerationResult
    monkeypatch.setattr(am, "CONFIG_DIR", tmp_path)

    async def fake_map(url, **kw):
        return _a_map()
    monkeypatch.setattr(am, "map_app", fake_map)
    seen = {}

    def generate(self, source, options=None, progress=None):
        seen["opts"] = options
        return GenerationResult(source=source, test_cases=[
            {"test_case_id": "T1", "title": "t", "description": "Covers: AC-1 | Technique: Use Case Testing."}])
    monkeypatch.setattr(gs.GenerationService, "generate", generate)
    body = {"url": "http://app.local", "story": "- Users can add a note\n- Empty notes are rejected", "project": "notes"}
    assert client.post("/api/functional/analyze", json=body).status_code == 400   # not authorized
    d = client.post("/api/functional/analyze", json=dict(body, authorized=True)).get_json()
    u = d["understanding"]
    assert [p["path"] for p in u["pages"]] == ["/login", "/notes"] and len(u["criteria"]) == 2
    assert seen["opts"].app_map is not None and "AC-2: Empty notes are rejected" in seen["opts"].repo_context
    assert d["coverage"]["uncovered"] == ["AC-2"] and d["coverage"]["percent"] == 50
    # a repository URL carrying credentials is refused before anything is cloned
    d = client.post("/api/functional/analyze", json={"repo": "https://user:tok@github.com/o/r", "story": "x y z"}).get_json()
    assert any("credentials" in n for n in d["understanding"]["notes"])
    assert client.post("/api/functional/analyze", json={}).status_code == 400


def test_web_functional_bugs_files_only_failures_and_stays_dry_run(client, monkeypatch):
    import core.jira_bugs as jb
    calls = {}
    real = jb.JiraBugFiler.__init__

    def init(self, *a, **kw):
        calls["dry_run"] = kw.get("dry_run")
        real(self, *a, **kw)
    monkeypatch.setattr(jb.JiraBugFiler, "__init__", init)
    res = [{"test_case_id": "T1", "status": "FAIL", "steps": [{"step_number": 1, "status": "FAIL", "actual_result": "no"}]},
           {"test_case_id": "T2", "status": "PASS"}]
    d = client.post("/api/functional/bugs", json={"results": res, "dry_run": False, "url": "http://a"}).get_json()
    assert d["dry_run"] is True and calls["dry_run"] is True          # live export is off on the server
    assert [b["test_case_id"] for b in d["bugs"]] == ["T1"]
    assert client.post("/api/functional/bugs", json={"results": res[1:]}).status_code == 400


def test_one_story_drives_functional_and_security_cases():
    from core.istqb_knowledge import prompt_block
    from core.traceability import extract_criteria, security_block, security_risks
    ac = extract_criteria("- A user can delete their own note\n- Admins can approve posts\n- Users can reset their password")
    risks = {(r["criterion"], r["risk"]) for r in security_risks(ac)}
    assert ("AC-1", "access control") in risks and ("AC-2", "function-level authorization") in risks
    assert ("AC-3", "authentication") in risks
    assert "AC-1 — access control" in security_block(security_risks(ac)) and security_block([]) == ""
    assert "negative access-control case" in prompt_block()


def test_requirement_verdicts_and_the_story_comment():
    from core.traceability import requirement_verdicts, verdict_comment
    ac = [{"id": "AC-1", "text": "add"}, {"id": "AC-2", "text": "reject"}, {"id": "AC-3", "text": "x"},
          {"id": "AC-4", "text": "y"}]
    cov = {"by_criterion": {"AC-1": ["T1", "T2"], "AC-2": ["T3"], "AC-3": [], "AC-4": ["T4"]}}
    res = [{"test_case_id": "T1", "status": "PASS"}, {"test_case_id": "T2", "status": "PASS"},
           {"test_case_id": "T3", "status": "FAIL"}, {"test_case_id": "T4", "status": "SKIP"}]
    v = {x["id"]: x["verdict"] for x in requirement_verdicts(ac, cov, res)}
    assert v == {"AC-1": "PASSED", "AC-2": "FAILED", "AC-3": "NOT COVERED", "AC-4": "NOT FULLY RUN"}
    text = verdict_comment(requirement_verdicts(ac, cov, res), "http://a", ["BUG-1"])
    assert "AC-2 FAILED" in text and "Bugs: BUG-1" in text


def test_bugs_link_to_the_story_and_the_story_gets_the_verdicts(client):
    res = [{"test_case_id": "T1", "status": "FAIL", "steps": [{"step_number": 1, "status": "FAIL", "actual_result": "no"}]}]
    body = {"results": res, "story_key": "notes-7", "url": "http://a",
            "criteria": [{"id": "AC-1", "text": "add a note"}],
            "coverage": {"by_criterion": {"AC-1": ["T1"]}}, "all_results": res}
    d = client.post("/api/functional/bugs", json=body).get_json()
    assert d["dry_run"] is True and d["results"][0]["linked_to"] == "NOTES-7"
    assert d["story_comment"]["story_key"] == "NOTES-7" and "AC-1 FAILED" in d["story_comment"]["text"]
    assert "regression_suite" not in d                                       # dry run writes nothing
    assert client.post("/api/functional/bugs", json=dict(body, story_key="not a key")).status_code == 400


def test_run_history_finds_regressions_fixes_and_flaky_cases(tmp_path):
    from core.run_history import add_regression_cases, compare, history, save_run
    for verdicts in ({"A": "PASS", "B": "FAIL", "C": "PASS"}, {"A": "PASS", "B": "FAIL", "C": "FAIL"},
                     {"A": "PASS", "B": "FAIL", "C": "PASS"}):
        save_run("p", "http://a", [{"test_case_id": k, "status": v} for k, v in verdicts.items()], root=tmp_path)
        import time
        time.sleep(0.002)
    now = [{"test_case_id": "A", "status": "FAIL"}, {"test_case_id": "B", "status": "PASS"},
           {"test_case_id": "C", "status": "FAIL"}, {"test_case_id": "D", "status": "PASS"}]
    ch = compare(history("p", root=tmp_path), now)
    assert ch["regressions"] == ["A", "C"] and ch["fixed"] == ["B"] and ch["new"] == ["D"]
    assert ch["flaky"] == ["C"] and ch["runs_before"] == 3
    path = add_regression_cases("p", [{"test_case_id": "A", "title": "t", "_verdict": {}}], {"A": "BUG-9"}, root=tmp_path)
    suite = json.loads(open(path).read())
    assert suite["test_cases"] == [{"test_case_id": "A", "title": "t", "regression_for": "BUG-9"}]


def test_a_bug_can_be_written_in_arabic(client):
    from core.jira_bugs import functional_result_to_bug
    r = {"test_case_id": "T1", "status": "FAIL", "steps": [{"step_number": 2, "status": "FAIL", "actual_result": "x"}]}
    bug = functional_result_to_bug(r, {"title": "Login"}, base_url="http://a", lang="ar")
    assert "يفشل عند الخطوة 2" in bug["description"] and "النتيجة الفعلية: x" in bug["description"]
    d = client.post("/api/functional/bugs", json={"results": [r], "lang": "ar"}).get_json()
    assert "يفشل عند الخطوة" in d["bugs"][0]["text"]
    assert "fails at step 2" in functional_result_to_bug(r, {"title": "Login"})["description"]


def test_doctor_names_what_is_missing_and_never_prints_secrets(monkeypatch, capsys):
    from core.doctor import print_report, run_checks
    monkeypatch.setenv("ANTHROPIC_API_KEY", "sk-ant-REAL-SECRET-VALUE")
    monkeypatch.setenv("JIRA_API_TOKEN", "your-api-token")                 # a placeholder is "not set"
    checks = {c["name"]: c for c in run_checks()}
    assert checks["AI provider"]["status"] == "ok" and checks["Jira"]["status"] == "warn"
    assert {"Python", "Packages", "Browser engine", "git", "Output folder", "Xray", "Live writes"} <= set(checks)
    print_report(list(checks.values()))
    assert "REAL-SECRET" not in capsys.readouterr().out


def test_reading_the_app_again_keeps_the_names_existing_cases_use(tmp_path, monkeypatch):
    import core.app_mapper as am
    from core.app_mapper import AppMap, save_map
    monkeypatch.setattr(am, "CONFIG_DIR", tmp_path)
    (tmp_path / "hand").mkdir()
    (tmp_path / "hand" / "selectors.json").write_text(json.dumps(
        {"_meta": {"auto_discovered": False}, "dashboard": {"note_input": "#note-input", "welcome_heading": "#welcome"}}))
    fresh = AppMap(base_url="http://a", pages=[{"path": "/dashboard", "elements": {"note": "#note-input", "add": "#add"}}])
    save_map(fresh, "hand")
    d = json.loads((tmp_path / "hand" / "selectors.json").read_text())["dashboard"]
    assert d == {"note": "#note-input", "add": "#add", "note_input": "#note-input", "welcome_heading": "#welcome"}
    # an auto-discovered file: aliases of live selectors survive, stale selectors go
    (tmp_path / "auto").mkdir()
    (tmp_path / "auto" / "selectors.json").write_text(json.dumps(
        {"_meta": {"auto_discovered": True, "manual_overrides": ["dashboard.pinned"]},
         "dashboard": {"note_input": "#note-input", "gone": "#old", "pinned": "#pin"}}))
    save_map(fresh, "auto")
    d = json.loads((tmp_path / "auto" / "selectors.json").read_text())["dashboard"]
    assert d == {"note": "#note-input", "add": "#add", "note_input": "#note-input", "pinned": "#pin"}


def test_an_automated_case_bug_states_the_expected_result_from_its_assertion():
    from core.jira_bugs import functional_result_to_bug
    case = {"title": "Login", "automation_steps": [{"playwright_action": "navigate", "value": "/"},
            {"playwright_action": "assert", "selector_key": "welcome_heading", "assertion_type": "text",
             "assertion_value": "Welcome, admin"}]}
    r = {"test_case_id": "T", "status": "FAIL", "steps": [{"step_number": 1, "status": "PASS"},
         {"step_number": 2, "status": "FAIL", "actual_result": "got 'Welcome, demo'"}]}
    assert "Expected: text check on welcome_heading: Welcome, admin" in functional_result_to_bug(r, case)["description"]


def test_the_session_route_says_what_is_saved_and_never_returns_it(client, tmp_path, monkeypatch):
    import core.session_state as ss
    monkeypatch.setattr(ss, "CONFIG_DIR", tmp_path / "projects")
    p = ss.state_path("shop")
    p.parent.mkdir(parents=True)
    p.write_text('{"cookies": [{"name": "sid", "value": "SECRET-SESSION"}]}')
    d = client.post("/api/session", json={"project": "shop"}).get_json()
    assert d["saved"] is True and "SECRET-SESSION" not in json.dumps(d)
    assert client.post("/api/session", json={"project": "shop", "action": "record", "url": "http://a"}).status_code == 400
    assert client.post("/api/session", json={}).status_code == 400
    assert client.post("/api/session", json={"project": "shop", "action": "forget"}).get_json()["forgotten"] is True
    assert not p.exists()


def test_a_second_ai_provider_speaks_the_same_client_shape(monkeypatch):
    """No Anthropic credit must not mean no test cases: QAX_AI_PROVIDER=openai sends the same
    request to any OpenAI-compatible API (OpenAI, Azure, a local Ollama) and answers in the shape
    every caller reads."""
    import threading

    from flask import Flask, request as freq
    from werkzeug.serving import make_server

    from core.ai_provider import OpenAICompatClient, make_ai_client
    seen = {}
    fake = Flask("fake-openai")

    @fake.post("/v1/chat/completions")
    def chat():
        seen["body"], seen["auth"] = freq.get_json(), freq.headers.get("Authorization")
        if seen["body"]["model"] == "broken":
            return {"error": "no such model"}, 404
        return {"model": "local-x", "choices": [{"message": {"content": '[{"title": "t"}]'}, "finish_reason": "length"}],
                "usage": {"prompt_tokens": 11, "completion_tokens": 7}}
    srv = make_server("127.0.0.1", 0, fake)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    try:
        monkeypatch.setenv("QAX_AI_PROVIDER", "openai")
        monkeypatch.setenv("QAX_AI_BASE_URL", f"http://127.0.0.1:{srv.server_port}/v1")
        monkeypatch.setenv("QAX_AI_MODEL", "local-x")
        monkeypatch.setenv("QAX_AI_API_KEY", "k-123")
        client = make_ai_client("ignored")
        assert isinstance(client, OpenAICompatClient)
        r = client.messages.create(model="claude-sonnet-4-5", max_tokens=100, system="be a QA engineer",
                                   messages=[{"role": "user", "content": [{"type": "text", "text": "story"}]},
                                             {"role": "assistant", "content": "["}])
        assert r.content[0].text == '[{"title": "t"}]' and r.stop_reason == "max_tokens"
        assert r.usage.input_tokens == 11 and r.usage.output_tokens == 7
        b = seen["body"]
        assert b["model"] == "local-x" and b["messages"][0] == {"role": "system", "content": "be a QA engineer"}
        assert b["messages"][-1]["role"] == "user" and "Start your answer with: [" in b["messages"][-1]["content"]
        assert seen["auth"] == "Bearer k-123"
        monkeypatch.setenv("QAX_AI_MODEL", "broken")
        with pytest.raises(RuntimeError, match="404"):
            make_ai_client().messages.create(messages=[{"role": "user", "content": "x"}])
        from services.context import QAXContext
        monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
        assert QAXContext.from_env().require_anthropic_key()                # the openai provider is enough
        monkeypatch.setenv("QAX_AI_PROVIDER", "anthropic")
        import anthropic
        assert isinstance(make_ai_client("sk-x"), anthropic.Anthropic)
    finally:
        srv.shutdown()


def test_the_package_metadata_matches_the_code():
    """pyproject.toml: same version as core.version, commands that exist, the runtime requirements."""
    import importlib
    import re as _re
    import tomllib
    from pathlib import Path as _P
    root = _P(__file__).resolve().parents[1]
    meta = tomllib.loads((root / "pyproject.toml").read_text())["project"]
    from core.version import __version__
    assert meta["version"] == __version__
    for target in meta["scripts"].values():
        mod, fn = target.split(":")
        assert callable(getattr(importlib.import_module(mod), fn)), target
    declared = {_re.split(r"[<>=!~\[]", d)[0].lower() for d in meta["dependencies"]}
    reqs = [l.split("#")[0].strip() for l in (root / "requirements.txt").read_text().splitlines()]
    runtime = {_re.split(r"[<>=!~\[]", r)[0].lower() for r in reqs if r and not r.startswith(("pytest",))}
    assert runtime <= declared, runtime - declared


def test_inputs_cannot_reach_secrets_or_leave_safe_channels(tmp_path, monkeypatch):
    """Security audit 2026-10-20: a repository folder's files go to the AI provider (never a hidden
    folder such as ~/.ssh); clones are https only; the MCP cases_file stays inside the workspace;
    the AI endpoint must be https unless it is this machine."""
    from core.ai_provider import OpenAICompatClient
    from core.mcp_server import call_tool
    from web.functional import _repo_source
    hidden = tmp_path / ".ssh"
    hidden.mkdir()
    assert "Hidden" in _repo_source(str(hidden))[1]
    assert "https" in _repo_source("http://example.com/x.git")[1]
    monkeypatch.setenv("QAX_REPO_ROOTS", str(tmp_path / "allowed"))
    (tmp_path / "other").mkdir()
    assert "outside" in _repo_source(str(tmp_path / "other"))[1]
    out = tmp_path / "outside.json"
    out.write_text("[]")
    monkeypatch.chdir(tmp_path / "other")
    d, err = call_tool("qax_run_tests", {"url": "http://a", "cases_file": str(out), "authorized": True})
    assert err and "inside the working folder" in d["error"]
    monkeypatch.setenv("QAX_AI_BASE_URL", "http://10.0.0.5/v1")
    with pytest.raises(RuntimeError, match="https"):
        OpenAICompatClient()
    monkeypatch.setenv("QAX_AI_BASE_URL", "http://127.0.0.1:11434/v1")
    assert OpenAICompatClient().base_url.startswith("http://127.0.0.1")


def test_zephyr_cloud_uses_the_v2_api_and_server_the_v1_path(monkeypatch):
    """Audit 2026-10-20: Cloud users were sent to the Server/DC path. Payload shape only —
    never exercised against a real Zephyr instance."""
    import requests as _rq

    from core.adapters.zephyr_adapter import ZephyrAdapter
    calls = []

    class R:
        ok, status_code = True, 201

        def json(self):
            return {"key": "QA-T1"}
    monkeypatch.setattr(_rq, "post", lambda url, **kw: calls.append((url, kw.get("json"))) or R())
    monkeypatch.setenv("JIRA_BASE_URL", "https://acme.atlassian.net")
    monkeypatch.setenv("ZEPHYR_API_TOKEN", "t")
    monkeypatch.setenv("ZEPHYR_PROJECT_KEY", "QA")
    tc = {"test_case_id": "T1", "title": "Login", "priority": "high",
          "manual_steps": [{"step_number": 1, "action": "open", "expected_result": "form"}]}
    ZephyrAdapter().export([tc], dry_run=False)
    assert calls[0][0] == "https://api.zephyrscale.smartbear.com/v2/testcases" and calls[0][1]["projectKey"] == "QA"
    assert calls[1][0].endswith("/testcases/QA-T1/teststeps") and calls[1][1]["items"][0]["inline"]["description"] == "open"
    calls.clear()
    monkeypatch.setenv("ZEPHYR_SERVER", "true")
    ZephyrAdapter().export([tc], dry_run=False)
    assert calls[0][0] == "https://acme.atlassian.net/rest/atm/1.0/testcase"


def test_excel_export_writes_a_workbook_people_can_open(tmp_path, monkeypatch):
    """Advertised on the landing page, never tested before the 2026-10-20 audit."""
    from openpyxl import load_workbook

    from core.adapters import get_adapter
    monkeypatch.chdir(tmp_path)
    tc = {"test_case_id": "TC-QA-0001", "title": "Login works", "priority": "high",
          "manual_steps": [{"step_number": 1, "action": "Open /login", "expected_result": "Form shown"}]}
    res = get_adapter("excel").export([tc], project="demo", dry_run=False)
    path = next(p for p in (res.get("references") or []) + [str(x) for x in tmp_path.rglob("*.xlsx")] if str(p).endswith(".xlsx"))
    wb = load_workbook(path)
    text = " ".join(str(c.value) for ws in wb for row in ws.iter_rows() for c in row if c.value)
    assert "Login works" in text and "Open /login" in text


def test_the_project_overview_shows_only_measured_data(client, tmp_path, monkeypatch):
    """Every card comes from a saved record; with no data a card is EMPTY and says how to get it —
    never a made-up number."""
    import core.app_mapper as am
    from core.project_overview import save_requirements
    from core.run_history import save_run
    monkeypatch.setattr(am, "CONFIG_DIR", tmp_path / "projects")
    d = client.get("/api/overview?project=shop").get_json()
    assert {c["id"] for c in d["cards"]} >= {"last_run", "changes", "requirements", "regression", "app_map", "security"}
    assert all(c["state"] == "empty" for c in d["cards"])
    save_requirements("shop", [{"id": "AC-1", "text": "add to basket"}, {"id": "AC-2", "text": "checkout"}],
                      {"by_criterion": {"AC-1": ["T1"], "AC-2": ["T2"]}})
    save_run("shop", "http://shop.local", [{"test_case_id": "T1", "status": "PASS"}, {"test_case_id": "T2", "status": "PASS"}])
    import time
    time.sleep(0.01)
    save_run("shop", "http://shop.local", [{"test_case_id": "T1", "status": "PASS"}, {"test_case_id": "T2", "status": "FAIL"}])
    cards = {c["id"]: c for c in client.get("/api/overview?project=shop").get_json()["cards"]}
    assert cards["last_run"]["value"] == "1 passed · 1 failed · 0 not run" and cards["last_run"]["state"] == "bad"
    assert "T2" in cards["changes"]["detail"] and cards["changes"]["value"].startswith("1 newly failing")
    assert cards["requirements"]["value"] == "1/2 criteria passed"
    assert [r["verdict"] for r in cards["requirements"]["rows"]] == ["PASSED", "FAILED"]
    assert cards["trend"]["points"] == [100, 50]
    assert client.get("/api/overview").status_code == 400


def test_one_unreadable_run_file_does_not_hide_the_history(tmp_path):
    from core.run_history import history, save_run
    save_run("p", "http://a", [{"test_case_id": "T", "status": "PASS"}], root=tmp_path)
    (tmp_path / "functional_runs" / "p" / "29991231T000000000000Z.json").write_text("{half-writ")
    assert [h["verdicts"] for h in history("p", root=tmp_path)] == [{"T": "PASS"}]


def test_overview_cards_for_gaps_flaky_critical_and_readiness(client, tmp_path, monkeypatch):
    """Coverage gaps, flaky cases (with quarantine), critical findings and release readiness —
    each computed from saved records; a quarantined case leaves bugs and readiness but still shows."""
    import time

    import core.app_mapper as am
    from core.project_overview import save_requirements
    from core.run_history import save_run
    monkeypatch.setattr(am, "CONFIG_DIR", tmp_path / "projects")
    save_requirements("shop", [{"id": "AC-1", "text": "pay"}, {"id": "AC-2", "text": "refund"}],
                      {"by_criterion": {"AC-1": ["T1", "T2"], "AC-2": []}, "uncovered": ["AC-2"]})
    for verdicts in (("PASS", "PASS"), ("PASS", "FAIL"), ("PASS", "PASS"), ("PASS", "FAIL")):
        save_run("shop", "http://shop.local", [{"test_case_id": "T1", "status": verdicts[0], "title": "pay ok"},
                                               {"test_case_id": "T2", "status": verdicts[1], "title": "pay twice",
                                                "cause": "app"}])
        time.sleep(0.01)
    cards = {c["id"]: c for c in client.get("/api/overview?project=shop").get_json()["cards"]}
    assert cards["coverage_gaps"]["rows"][0]["id"] == "AC-2"
    assert [r["id"] for r in cards["flaky"]["rows"]] == ["T2"] and cards["flaky"]["rows"][0]["verdict"] == "FLAKY"
    assert cards["critical"]["rows"][0]["id"] == "T2" and cards["readiness"]["value"] == "Not ready"
    assert client.post("/api/quarantine", json={"project": "shop", "case_id": "T2", "on": True}).status_code == 200
    cards = {c["id"]: c for c in client.get("/api/overview?project=shop").get_json()["cards"]}
    assert cards["flaky"]["rows"][0]["verdict"] == "QUARANTINED" and cards["critical"]["value"] == "None open"
    assert not any("failing" in r["text"] for r in cards["readiness"]["rows"])   # quarantined: not a blocker
    assert any("criteria not passed" in r["text"] for r in cards["readiness"]["rows"])  # AC-2 has no case
    assert client.post("/api/quarantine", json={"project": "shop"}).status_code == 400
