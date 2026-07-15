"""
core/adapters/excel_adapter.py — Excel Export Adapter
======================================================
Writes test cases to a colour-formatted .xlsx workbook using openpyxl.

The workbook has one sheet ("Test Cases") with purple header row and one
data row per test case. Multi-step fields (manual steps, expected results)
are joined with newlines so they read naturally in Excel.

Output path:
    output/test_cases_{project}_{YYYYMMDD_HHMMSS}.xlsx

Requires:
    pip install openpyxl>=3.1.0

Usage:
    adapter = ExcelAdapter()
    valid, err = adapter.validate_config()
    result = adapter.export(test_cases, project="myapp", dry_run=True)
"""

import logging
import os
from datetime import datetime
from pathlib import Path

from core.adapters.base import BaseExportAdapter

logger = logging.getLogger(__name__)

# ── Column definitions: (header_label, key_or_callable) ───────────────────────
_COLUMNS = [
    ("ID",             lambda tc, _: tc.get("test_case_id", "")),
    ("Title",          lambda tc, _: tc.get("title", "")),
    ("Module",         lambda tc, _: tc.get("module", "")),
    ("Type",           lambda tc, _: tc.get("test_type", "")),
    ("Priority",       lambda tc, _: tc.get("priority", "Medium")),
    ("Confidence",     lambda tc, _: tc.get("confidence_score", 0)),
    ("Preconditions",  lambda tc, _: tc.get("preconditions", "")),
    ("Manual Steps",   lambda tc, steps: "\n".join(
        f"{s.get('step_number', i+1)}. {s.get('action', '')}"
        for i, s in enumerate(steps)
    )),
    ("Expected Result", lambda tc, steps: "\n".join(
        s.get("expected_result", "") for s in steps
    )),
    ("Tags",           lambda tc, _: ", ".join(tc.get("tags", []))),
]

# Purple brand colour (matches QAX UI theme)
_HEADER_COLOR = "7C3AED"
_HEADER_FONT  = "FFFFFF"


class ExcelAdapter(BaseExportAdapter):
    """
    Export adapter that saves test cases as a formatted Excel workbook.

    Uses openpyxl directly — no pandas dependency required.
    Column widths are set to sensible defaults; users can adjust in Excel.
    In dry_run mode the file is NOT written; only the target path is returned.
    """

    name = "excel"

    def validate_config(self) -> tuple[bool, str]:
        """
        Check that openpyxl is installed.

        Returns:
            tuple[bool, str]: (True, "") if openpyxl is importable,
                              (False, install hint) otherwise.
        """
        try:
            import openpyxl  # noqa: F401
            return True, ""
        except ImportError:
            return False, "openpyxl not installed — run: pip install openpyxl>=3.1.0"

    def export(
        self,
        test_cases: list[dict],
        project: str | None = None,
        dry_run: bool = True,
    ) -> dict:
        """
        Write test cases to a .xlsx file.

        Args:
            test_cases (list[dict]): Validated QAX test case dicts.
            project (str | None): Used in the output filename.
                                   Falls back to DEFAULT_PROJECT env var.
            dry_run (bool): When True, log the target path without writing.

        Returns:
            dict: Standard result dict with "references" containing the file path.
        """
        if not test_cases:
            return self._empty_result(0, error="No test cases provided")

        valid, err = self.validate_config()
        if not valid:
            logger.error("QAX ExcelAdapter: %s", err)
            return self._empty_result(len(test_cases), error=err)

        proj      = project or os.getenv("DEFAULT_PROJECT", "export")
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_dir = Path("output")
        output_dir.mkdir(parents=True, exist_ok=True)
        filename  = str(output_dir / f"test_cases_{proj}_{timestamp}.xlsx")

        if dry_run:
            logger.info(
                "QAX ExcelAdapter [DRY RUN]: would save %d case(s) to %s",
                len(test_cases), filename,
            )
            print(f"  [DRY RUN] Would save {len(test_cases)} test case(s) to: {filename}")
            return self._success_result(len(test_cases), 0, [filename])

        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Test Cases"

            # ── Header row ────────────────────────────────────────────────────
            header_fill = PatternFill(
                start_color=_HEADER_COLOR,
                end_color=_HEADER_COLOR,
                fill_type="solid",
            )
            header_font = Font(bold=True, color=_HEADER_FONT)
            col_widths  = [12, 45, 18, 15, 12, 12, 30, 50, 45, 25]

            for col_idx, (label, _) in enumerate(_COLUMNS, 1):
                cell = ws.cell(row=1, column=col_idx, value=label)
                cell.font      = header_font
                cell.fill      = header_fill
                cell.alignment = Alignment(vertical="center", wrap_text=False)
                col_letter     = ws.cell(row=1, column=col_idx).column_letter
                ws.column_dimensions[col_letter].width = col_widths[col_idx - 1]

            # Freeze header row
            ws.freeze_panes = "A2"

            # ── Data rows ─────────────────────────────────────────────────────
            wrap = Alignment(vertical="top", wrap_text=True)
            for row_idx, tc in enumerate(test_cases, 2):
                steps = tc.get("manual_steps", [])
                for col_idx, (_, extractor) in enumerate(_COLUMNS, 1):
                    value = extractor(tc, steps)
                    cell  = ws.cell(row=row_idx, column=col_idx, value=value)
                    cell.alignment = wrap
                ws.row_dimensions[row_idx].height = max(
                    15, min(15 * len(steps), 120)
                )

            wb.save(filename)
            logger.info(
                "QAX ExcelAdapter: saved %d test case(s) → %s",
                len(test_cases), filename,
            )
            return self._success_result(len(test_cases), 0, [filename])

        except Exception as exc:
            logger.error("QAX ExcelAdapter: write failed — %s", exc)
            return self._empty_result(len(test_cases), error=str(exc))
